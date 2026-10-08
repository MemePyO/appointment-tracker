"""Backend: data model, day parsing, project save/load, Excel export."""

import json
import os
import re

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

APP_NAME = "Appointment Tracker"

GREEN_HEADER = "E2EFDA"
MONTH_BLUE = "B8CCE4"
DAY_GREEN = "92D050"
DATA_BLUE = "D9E1F2"
YELLOW_PALE = "FFFFE0"
YELLOW_BRIGHT = "FFFF00"
SPACER_BLUE = "BDD7EE"
TOTAL_GREEN = "C6E0B4"


def default_patients_file():
    return os.path.join(os.path.dirname(os.path.abspath(__file__)), "patients.json")


def _unique_names(names):
    unique = []
    seen = set()
    for item in names:
        name = str(item).strip()
        if name and name.lower() not in seen:
            seen.add(name.lower())
            unique.append(name)
    return unique


def load_patients_list(path=None):
    path = path or default_patients_file()
    if not os.path.exists(path):
        return []
    try:
        with open(path, encoding="utf-8") as fh:
            data = json.load(fh)
    except json.JSONDecodeError as exc:
        raise ValueError(f"Not a valid JSON file: {exc}") from exc
    if not isinstance(data, list):
        raise ValueError("Patients list must be a JSON array of names.")
    return _unique_names(data)


def save_patients_list(names, path=None):
    path = path or default_patients_file()
    unique = _unique_names(names)
    with open(path, "w", encoding="utf-8") as fh:
        json.dump(unique, fh, ensure_ascii=False, indent=2)
    return unique


def add_patient_name(name, path=None):
    name = (name or "").strip()
    current = load_patients_list(path)
    if name and name.lower() not in {n.lower() for n in current}:
        current.append(name)
        save_patients_list(current, path)
    return current


def filter_names(names, prefix):
    prefix = (prefix or "").strip().lower()
    if not prefix:
        return list(names)
    starts = [n for n in names if n.lower().startswith(prefix)]
    contains = [n for n in names
                if not n.lower().startswith(prefix) and prefix in n.lower()]
    return starts + contains


def parse_days(text, max_day=None):
    days = set()
    text = (text or "").strip()
    if not text:
        return []
    for part in text.split(","):
        part = part.strip()
        if not part:
            continue
        m = re.fullmatch(r"(\d+)\s*-\s*(\d+)", part)
        if m:
            a, b = int(m.group(1)), int(m.group(2))
            if a > b:
                raise ValueError(f"Invalid range '{part}': start is bigger than end.")
            if b - a > 366:
                raise ValueError(f"Range too long: '{part}'.")
            days.update(range(a, b + 1))
        elif part.isdigit():
            days.add(int(part))
        else:
            raise ValueError(f"Invalid day value: '{part}'.")
    if days:
        low, high = min(days), max(days)
        if low < 1 or (max_day is not None and high > max_day):
            limit = max_day if max_day is not None else high
            raise ValueError(f"Days must be between 1 and {limit}.")
    return sorted(days)


def days_to_text(days):
    return ", ".join(str(d) for d in days)


def new_project():
    return {
        "title": "",
        "month": "",
        "days_in_month": 30,
        "c_days": [],
        "l_days": [],
        "patients": [],
    }


def _days_list_to_text(value):
    if isinstance(value, str):
        return value
    if isinstance(value, list):
        return days_to_text(value)
    raise ValueError("Days must be a list of numbers or a text like '1, 3, 5-7'.")


def validate_project(project):
    if not isinstance(project, dict):
        raise ValueError("Project must be a JSON object.")
    for key in ("title", "month", "days_in_month", "c_days", "l_days", "patients"):
        if key not in project:
            raise ValueError(f"Missing field: {key}")
    n = project["days_in_month"]
    if not isinstance(n, int) or not 1 <= n <= 31:
        raise ValueError("Days must be a whole number between 1 and 31.")
    if not isinstance(project["c_days"], list) or not isinstance(project["l_days"], list):
        raise ValueError("C days and L days must be lists.")
    if not isinstance(project["patients"], list):
        raise ValueError("Patients must be a list.")
    patients = []
    for p in project["patients"]:
        if not isinstance(p, dict):
            raise ValueError("Each patient must be an object.")
        x_days = parse_days(_days_list_to_text(p.get("x_days", [])), n)
        patients.append({
            "name": str(p.get("name", "")),
            "x_days": x_days,
            "decision_date": str(p.get("decision_date", "")),
            "nr_zile": str(p.get("nr_zile", "")),
            "localitate": str(p.get("localitate", "")),
        })
    return {
        "title": str(project["title"]),
        "month": str(project["month"]),
        "days_in_month": n,
        "c_days": parse_days(_days_list_to_text(project["c_days"]), n),
        "l_days": parse_days(_days_list_to_text(project["l_days"]), n),
        "patients": patients,
    }


def save_project(project, path):
    with open(path, "w", encoding="utf-8") as fh:
        json.dump(project, fh, ensure_ascii=False, indent=2)


def load_project(path):
    try:
        with open(path, encoding="utf-8") as fh:
            project = json.load(fh)
    except json.JSONDecodeError as exc:
        raise ValueError(f"Not a valid JSON file: {exc}") from exc
    return validate_project(project)


def _number(value):
    text = str(value).strip()
    if text.isdigit():
        return int(text)
    return text or None


def export_excel(project, path):
    project = validate_project(project)
    n = project["days_in_month"]
    c_set = set(project["c_days"])
    l_set = set(project["l_days"])
    patients = project["patients"]
    if not patients:
        raise ValueError("Add at least one patient before exporting.")

    first_day = 3
    last_day = 2 + n
    count_col = last_day + 1
    dec_col = count_col + 1
    zile_col = dec_col + 1
    loc_col = zile_col + 1

    wb = Workbook()
    ws = wb.active
    ws.title = "TABEL"

    thin = Side(style="thin", color="000000")
    box = Border(left=thin, right=thin, top=thin, bottom=thin)

    def fill(hexcode):
        return PatternFill("solid", fgColor=hexcode)

    f_header = fill(GREEN_HEADER)
    f_month = fill(MONTH_BLUE)
    f_dayhdr = fill(DAY_GREEN)
    f_data = fill(DATA_BLUE)
    f_yellow = fill(YELLOW_PALE)
    f_bright = fill(YELLOW_BRIGHT)
    f_blue = fill(SPACER_BLUE)
    f_total = fill(TOTAL_GREEN)

    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left = Alignment(horizontal="left", vertical="center")
    bold = Font(bold=True)

    def style_range(r1, c1, r2, c2, fill_obj=None, border=None, font=None,
                    align=None, value=None):
        if value is not None:
            ws.cell(row=r1, column=c1, value=value)
        for r in range(r1, r2 + 1):
            for c in range(c1, c2 + 1):
                cell = ws.cell(row=r, column=c)
                if fill_obj is not None:
                    cell.fill = fill_obj
                if border is not None:
                    cell.border = border
                if font is not None:
                    cell.font = font
                if align is not None:
                    cell.alignment = align

    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=loc_col)
    title = ws.cell(row=1, column=1, value=project["title"])
    title.font = Font(bold=True, size=14)
    title.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 30

    ws.merge_cells(start_row=2, start_column=1, end_row=3, end_column=1)
    ws.merge_cells(start_row=2, start_column=2, end_row=3, end_column=2)
    style_range(2, 1, 3, 1, f_header, box, bold, center, value="NR.")
    style_range(2, 2, 3, 2, f_header, box, bold,
                Alignment(horizontal="left", vertical="center"), value="PACIENT")

    ws.merge_cells(start_row=2, start_column=first_day, end_row=2, end_column=last_day)
    style_range(2, first_day, 2, last_day, f_month, box, bold, center,
                value=project["month"])
    style_range(3, first_day, 3, last_day, f_dayhdr, box, bold, center)
    for d in range(1, n + 1):
        ws.cell(row=3, column=first_day + d - 1, value=d)

    ws.merge_cells(start_row=2, start_column=count_col, end_row=3, end_column=count_col)
    style_range(2, count_col, 3, count_col, f_yellow, box, bold, center)

    for col, text in ((dec_col, "DATA DECIZIE"), (zile_col, "NR ZILE"),
                      (loc_col, "LOCALITATE")):
        ws.merge_cells(start_row=2, start_column=col, end_row=3, end_column=col)
        style_range(2, col, 3, col, f_bright, box, bold, center, value=text)

    ws.row_dimensions[2].height = 24
    ws.row_dimensions[3].height = 20
    ws.row_dimensions[4].height = 8

    row = 5
    first_data_row = row
    for i, p in enumerate(patients, start=1):
        ws.cell(row=row, column=1, value=i)
        ws.cell(row=row, column=2, value=p["name"])
        x_set = set(p["x_days"])
        for d in range(1, n + 1):
            if d in x_set:
                value = "X"
            elif d in c_set:
                value = "C"
            elif d in l_set:
                value = "L"
            else:
                value = None
            ws.cell(row=row, column=first_day + d - 1, value=value)
        ws.cell(
            row=row,
            column=count_col,
            value=(
                f'=COUNTIF({get_column_letter(first_day)}{row}:'
                f'{get_column_letter(last_day)}{row},"X")'
            ),
        )
        ws.cell(row=row, column=dec_col, value=p["decision_date"] or None)
        ws.cell(row=row, column=zile_col, value=_number(p["nr_zile"]))
        ws.cell(row=row, column=loc_col, value=p["localitate"] or None)
        for c in range(1, loc_col + 1):
            cell = ws.cell(row=row, column=c)
            cell.border = box
            if c <= 2:
                cell.fill = f_header
            elif first_day <= c <= last_day:
                cell.fill = f_data
            else:
                cell.fill = f_yellow
            cell.alignment = left if c in (2, loc_col) else center
        ws.row_dimensions[row].height = 18
        row += 1
    last_data_row = row - 1

    style_range(row, 1, row, loc_col, f_blue, box)
    ws.row_dimensions[row].height = 8

    total_row = row + 1
    ws.cell(row=total_row, column=2, value="TOTAL")
    ws.cell(
        row=total_row,
        column=count_col,
        value=(
            f"=SUM({get_column_letter(count_col)}{first_data_row}:"
            f"{get_column_letter(count_col)}{last_data_row})"
        ),
    )
    for c in range(1, loc_col + 1):
        cell = ws.cell(row=total_row, column=c)
        cell.border = box
        cell.fill = f_total if c <= count_col else f_yellow
        cell.alignment = left if c == 2 else center
        if cell.value is not None:
            cell.font = bold
    ws.row_dimensions[total_row].height = 18

    style_range(total_row + 1, 1, total_row + 1, loc_col, f_blue, box)
    ws.row_dimensions[total_row + 1].height = 8

    ws.column_dimensions["A"].width = 5
    ws.column_dimensions["B"].width = 24
    for c in range(first_day, last_day + 1):
        ws.column_dimensions[get_column_letter(c)].width = 3.6
    ws.column_dimensions[get_column_letter(count_col)].width = 6.5
    ws.column_dimensions[get_column_letter(dec_col)].width = 13
    ws.column_dimensions[get_column_letter(zile_col)].width = 9
    ws.column_dimensions[get_column_letter(loc_col)].width = 20

    wb.save(path)
    return path
