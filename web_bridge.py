"""JSON-string API over backend.py for the web frontend (runs under Pyodide
and under desktop Python for tests)."""

import base64
import json
import os
import sys
import tempfile

from openpyxl import load_workbook

import appointment
import backend
import lang
import weekend
import workbook

MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _temp_path():
    if sys.platform == "emscripten":
        return "/tmp/appointments_work.xlsx"
    return os.path.join(tempfile.gettempdir(), "appointments_bridge.xlsx")


def _read_b64(path):
    with open(path, "rb") as fh:
        data = fh.read()
    try:
        os.remove(path)
    except OSError:
        pass
    return base64.b64encode(data).decode("ascii")


def _write_b64(path, b64text):
    with open(path, "wb") as fh:
        fh.write(base64.b64decode(b64text))


def parse_days_json(text, max_day=None):
    try:
        days = backend.parse_days(text, max_day if max_day is None else int(max_day))
        return json.dumps({"ok": True, "days": days})
    except ValueError as exc:
        return json.dumps({"ok": False, "error": str(exc)})


def filter_names_json(names_json, prefix):
    names = json.loads(names_json)
    return json.dumps(backend.filter_names(names, prefix))


def validate_project_json(project_json):
    try:
        project = backend.validate_project(json.loads(project_json))
        return json.dumps({"ok": True, "project": project})
    except ValueError as exc:
        return json.dumps({"ok": False, "error": str(exc)})


def export_xlsx_b64(project_json):
    try:
        project = json.loads(project_json)
    except json.JSONDecodeError as exc:
        return json.dumps({"ok": False, "error": f"Invalid JSON: {exc}"})
    path = _temp_path()
    try:
        backend.export_excel(project, path)
    except ValueError as exc:
        return json.dumps({"ok": False, "error": str(exc)})
    return json.dumps({
        "ok": True,
        "filename": "appointments.xlsx",
        "mime": MIME,
        "b64": _read_b64(path),
    })


def month_sheet_name_json(month, year=None):
    try:
        return json.dumps({"ok": True,
                           "sheet": workbook.month_sheet_name(month, year)})
    except ValueError as exc:
        return json.dumps({"ok": False, "error": str(exc)})


def existing_sheets_b64(workbook_b64):
    path = _temp_path()
    try:
        _write_b64(path, workbook_b64)
        wb = load_workbook(path)
        return json.dumps({"ok": True, "sheets": wb.sheetnames})
    except Exception as exc:
        return json.dumps({"ok": False, "error": str(exc)})
    finally:
        try:
            os.remove(path)
        except OSError:
            pass


def export_workbook_b64(project_json, workbook_b64=""):
    try:
        project = json.loads(project_json)
    except json.JSONDecodeError as exc:
        return json.dumps({"ok": False, "error": f"Invalid JSON: {exc}"})
    path = _temp_path()
    try:
        if workbook_b64:
            _write_b64(path, workbook_b64)
        path, sheet = workbook.export_to_workbook(project, path)
    except ValueError as exc:
        return json.dumps({"ok": False, "error": str(exc)})
    return json.dumps({
        "ok": True,
        "b64": _read_b64(path),
        "sheet": sheet,
        "mime": MIME,
        "filename": "appointments.xlsx",
    })


def add_appointment_b64(workbook_b64, name, day, month, year):
    path = _temp_path()
    try:
        _write_b64(path, workbook_b64)
        path, sheet, result = appointment.add_appointment(
            path, name, day, month, year)
    except ValueError as exc:
        return json.dumps({"ok": False, "error": str(exc)})
    except Exception as exc:
        return json.dumps({"ok": False, "error": f"Failed: {exc}"})
    return json.dumps({
        "ok": True,
        "b64": _read_b64(path),
        "sheet": sheet,
        "result": result,
    })


def set_language_json(code):
    lang.set_language(code)
    return json.dumps({"ok": True, "lang": lang.get_language()})


def strings_json(code=None):
    return lang.strings_json(code)


def weekend_days_json(month, year=None):
    try:
        return json.dumps({"ok": True,
                           "days": weekend.weekend_days(month, year)})
    except ValueError as exc:
        return json.dumps({"ok": False, "error": str(exc)})
