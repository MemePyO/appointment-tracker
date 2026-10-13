"""Introdu sedinta: mark one appointment (X) by date in a month workbook.

Called from frontend (tab "Introdu ședință"); backend.py / workbook.py reused,
not changed.
"""

import os
import re
from calendar import monthrange
from copy import copy

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, PatternFill, Side
from openpyxl.utils import get_column_letter

import backend
from workbook import month_sheet_name

FIRST_DAY_COL = 3
FIRST_DATA_ROW = 5
NAME_COL = 2

GREEN_HEADER = "E2EFDA"
DATA_BLUE = "B8CCE4"
YELLOW_PALE = "FFFFE0"


def _open(path):
    if not os.path.exists(path):
        wb = Workbook()
        wb.remove(wb.active)
        return wb
    try:
        return load_workbook(path)
    except Exception as exc:
        raise ValueError(f"Not a valid Excel file: {path}\n{exc}")


def _find_row(ws, name):
    target = str(name).strip().casefold()
    for r in range(FIRST_DATA_ROW, ws.max_row + 1):
        value = ws.cell(row=r, column=NAME_COL).value
        if value is None or not str(value).strip():
            continue
        text = str(value).strip()
        if text == "TOTAL":
            break
        if text.casefold() == target:
            return r
    return None


def _total_row(ws):
    for r in range(FIRST_DATA_ROW, ws.max_row + 1):
        if str(ws.cell(row=r, column=NAME_COL).value or "").strip() == "TOTAL":
            return r
    return ws.max_row + 2


def _count_col(ws):
    for c in range(FIRST_DAY_COL, (ws.max_column or FIRST_DAY_COL) + 1):
        if ws.cell(row=2, column=c).value == "COUNT":
            return c
    c = FIRST_DAY_COL
    while isinstance(ws.cell(row=3, column=c).value, int):
        c += 1
    return c


def _parse_date(text):
    parts = [p for p in re.split(r"[./\-\s]+", str(text or "").strip()) if p]
    if len(parts) != 3:
        raise ValueError("Date must look like 30.04.2026")
    day, month, year = parts
    return day, month, year


def add_appointment(path, name, day, month, year=None):
    name = str(name or "").strip()
    if not name:
        raise ValueError("Patient name is empty.")
    try:
        day = int(str(day).strip())
    except ValueError:
        raise ValueError("Day must be a whole number like 30.")

    sheet_name = month_sheet_name(month, year)
    num, yr = (int(p) for p in sheet_name.split("."))
    last_day = monthrange(yr, num)[1]
    if not 1 <= day <= last_day:
        raise ValueError(f"Day must be between 1 and {last_day} for {sheet_name}.")

    wb = _open(path)
    if sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
    else:
        ws = wb.create_sheet(sheet_name)
        backend._write_table(ws, {
            "title": "",
            "month": sheet_name,
            "year": yr,
            "days_in_month": last_day,
            "c_days": [],
            "l_days": [],
            "patients": [],
        })

    count_col = _count_col(ws)
    loc_col = count_col + 3
    day_col = FIRST_DAY_COL + day - 1

    row = _find_row(ws, name)
    added = False
    if row is None:
        total = _total_row(ws)
        row = max(FIRST_DATA_ROW, total - 1)
        ws.insert_rows(row)
        src = row - 1
        for c in range(1, loc_col + 1):
            cell = ws.cell(row=row, column=c)
            if src >= FIRST_DATA_ROW:
                cell._style = copy(ws.cell(row=src, column=c)._style)
            else:
                thin = Side(style="thin", color="000000")
                cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
                if c <= 2:
                    cell.fill = PatternFill("solid", fgColor=GREEN_HEADER)
                elif FIRST_DAY_COL <= c <= FIRST_DAY_COL + last_day - 1:
                    cell.fill = PatternFill("solid", fgColor=DATA_BLUE)
                else:
                    cell.fill = PatternFill("solid", fgColor=YELLOW_PALE)
                cell.alignment = (Alignment(horizontal="left", vertical="center")
                                  if c in (2, loc_col) else
                                  Alignment(horizontal="center", vertical="center"))
        ws.row_dimensions[row].height = 18
        ws.cell(row=row, column=NAME_COL, value=name)
        ws.cell(row=row, column=count_col,
                value=f'=COUNTIF({get_column_letter(FIRST_DAY_COL)}{row}:'
                      f'{get_column_letter(FIRST_DAY_COL + last_day - 1)}{row},"X")')
        added = True

        total = _total_row(ws)
        for r in range(FIRST_DATA_ROW, total):
            if ws.cell(row=r, column=NAME_COL).value:
                ws.cell(row=r, column=1, value=r - FIRST_DATA_ROW + 1)
        ws.cell(row=total, column=count_col,
                value=f"=SUM({get_column_letter(count_col)}{FIRST_DATA_ROW}:"
                      f"{get_column_letter(count_col)}{total - 1})")

    cell = ws.cell(row=row, column=day_col)
    result = "added" if added and cell.value != "X" else "updated"
    if cell.value == "X":
        result = "already_marked"
    else:
        cell.value = "X"
    wb.save(path)
    return path, sheet_name, result


def build_tab(parent, get_workbook, entry_factory=None, on_path=None):
    import customtkinter as ctk
    from tkinter import filedialog

    frame = ctk.CTkFrame(parent, fg_color="transparent")
    frame.pack(fill="both", expand=True, padx=12, pady=12)

    ctk.CTkLabel(
        frame,
        text="Mark an appointment (X) for a patient on a date in the Excel "
             "workbook. The patient is added to the month sheet if missing.",
        text_color="#555555",
    ).pack(anchor="w")

    row1 = ctk.CTkFrame(frame, fg_color="transparent")
    row1.pack(fill="x", pady=(14, 4))

    ctk.CTkLabel(row1, text="Patient name:").pack(side="left", padx=(0, 6))
    if entry_factory is not None:
        entry_name = entry_factory(row1, width=280,
                                   placeholder_text="e.g. STOICA FLORIN")
    else:
        entry_name = ctk.CTkEntry(row1, width=280,
                                  placeholder_text="e.g. STOICA FLORIN")
    entry_name.pack(side="left", padx=(0, 16))

    ctk.CTkLabel(row1, text="Date:").pack(side="left", padx=(0, 6))
    entry_date = ctk.CTkEntry(row1, width=140, placeholder_text="30.04.2026")
    entry_date.pack(side="left")

    row2 = ctk.CTkFrame(frame, fg_color="transparent")
    row2.pack(fill="x", pady=(10, 4))

    status = ctk.CTkLabel(row2, text="", anchor="w", text_color="#555555")

    def _mark():
        try:
            day, month, year = _parse_date(entry_date.get())
        except ValueError as exc:
            status.configure(text=str(exc), text_color="#c0392b")
            return
        path = get_workbook()
        if not path:
            path = filedialog.askopenfilename(
                title="Select Excel workbook",
                filetypes=[("Excel file", "*.xlsx"), ("All files", "*.*")])
            if not path:
                status.configure(
                    text="No workbook selected — export one first.",
                    text_color="#c0392b")
                return
            if on_path is not None:
                on_path(path)
        try:
            path, sheet, result = add_appointment(
                path, entry_name.get(), day, month, year)
        except ValueError as exc:
            status.configure(text=str(exc), text_color="#c0392b")
            return
        except Exception as exc:
            status.configure(text=f"Failed: {exc}", text_color="#c0392b")
            return
        if on_path is not None:
            on_path(path)
        who = entry_name.get().strip()
        if result == "already_marked":
            status.configure(text=f"Already marked: {who} on {sheet}",
                             text_color="#555555")
        elif result == "added":
            status.configure(text=f"Added {who} + X on {sheet}",
                             text_color="#2e7d32")
        else:
            status.configure(text=f"Marked X for {who} on {sheet}",
                             text_color="#2e7d32")

    btn = ctk.CTkButton(row2, text="Mark appointment", width=160,
                        fg_color="#2e7d32", hover_color="#1b5e20",
                        command=_mark)
    btn.pack(side="left")
    status.pack(side="left", padx=12)

    frame.entry_name = entry_name
    frame.entry_date = entry_date
    frame.status = status
    frame.mark_now = _mark
    return frame
