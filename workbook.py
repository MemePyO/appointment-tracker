"""Multi-sheet Excel workbook export: one sheet per month, named MM.YYYY.

Uses backend._write_table for the table layout; backend.py itself is
otherwise unchanged.
"""

import os
from datetime import date

from openpyxl import Workbook, load_workbook

import backend

MONTHS = {
    "ianuarie": 1, "februarie": 2, "martie": 3, "aprilie": 4,
    "mai": 5, "iunie": 6, "iulie": 7, "august": 8,
    "septembrie": 9, "octombrie": 10, "noiembrie": 11, "decembrie": 12,
}


def _check_year(year):
    if year is None:
        return date.today().year
    try:
        year = int(str(year).strip())
    except ValueError:
        raise ValueError("Year must be a whole number like 2026.")
    if not 2000 <= year <= 2100:
        raise ValueError("Year must be between 2000 and 2100.")
    return year


def month_sheet_name(month, year=None):
    year = _check_year(year)
    text = str(month or "").strip().casefold()
    if not text:
        raise ValueError("Month is empty — use ianuarie…decembrie or a number like 9.")
    if text.isdigit():
        num = int(text)
        if not 1 <= num <= 12:
            raise ValueError(f"Invalid month number: {text}.")
    else:
        num = MONTHS.get(text)
        if num is None:
            raise ValueError(
                f"Unknown month '{month}' — use ianuarie…decembrie or a number like 9.")
    return f"{num:02d}.{year}"


def export_to_workbook(project, path):
    project = backend.validate_project(project)
    if not project["patients"]:
        raise ValueError("Add at least one patient before exporting.")
    sheet_name = month_sheet_name(project["month"], project.get("year"))

    if os.path.exists(path):
        try:
            wb = load_workbook(path)
        except Exception as exc:
            raise ValueError(f"Not a valid Excel file: {path}\n{exc}")
        if sheet_name in wb.sheetnames:
            index = wb.sheetnames.index(sheet_name)
            del wb[sheet_name]
            ws = wb.create_sheet(sheet_name, index)
        else:
            ws = wb.create_sheet(sheet_name)
    else:
        wb = Workbook()
        wb.remove(wb.active)
        ws = wb.create_sheet(sheet_name)

    backend._write_table(ws, project)
    wb.save(path)
    return path, sheet_name
